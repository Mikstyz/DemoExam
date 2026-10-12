"""Import the provided XLSX resources into the Chudo Obuv PostgreSQL database.

Run from the project root:
    python -m scripts.import_data
    python -m scripts.import_data --reset

The script imports current stock as-is. Historical orders do not decrement stock
again, and historical order-item prices are preserved from Orders_import.xlsx.
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlmodel import Session

BASE_DIR = Path(__file__).resolve().parents[1]
DEFAULT_INPUT_DIR = BASE_DIR / "import_data"

from db.database import create_tables, engine  # noqa: E402
from model import (
    Category,
    Order,
    OrderItem,
    Product,
    Role,
    Size,
    StockItem,
    Subcategory,
    User,
)  # noqa: E402
from repositories import (  # noqa: E402
    CategoryRepository,
    OrderItemRepository,
    OrderRepository,
    ProductRepository,
    RoleRepository,
    SizeRepository,
    StockItemRepository,
    SubcategoryRepository,
    UserRepository,
)

FILES = {
    "users": "Users_import.xlsx",
    "products": "Products_import.xlsx",
    "sizes": "Sizes_import.xlsx",
    "stock": "Stock_Items_import.xlsx",
    "orders": "Orders_import.xlsx",
}

ROLE_NAMES = {
    "администратор": "Администратор",
    "менеджер": "Менеджер",
    "авторизованный пользователь": "Авторизованный пользователь",
    "пользователь": "Авторизованный пользователь",
    "user": "Авторизованный пользователь",
    "admin": "Администратор",
    "manager": "Менеджер",
}


def clean_text(
    value: Any, *, required: bool = False, field: str = "значение"
) -> str | None:
    if value is None:
        if required:
            raise ValueError(f"Пустое обязательное поле: {field}")
        return None
    result = " ".join(str(value).strip().split())
    if not result:
        if required:
            raise ValueError(f"Пустое обязательное поле: {field}")
        return None
    return result


def normalize_key(value: Any) -> str:
    return (clean_text(value) or "").casefold()


def normalized_name(value: Any) -> str:
    """Normalize category/subcategory names for stable deduplication."""
    return " ".join((clean_text(value) or "").split()).casefold()


def normalize_size(value: Any) -> str:
    """Convert Excel values 20.0 -> '20', but retain half sizes like 38.5."""
    if value is None:
        raise ValueError("В файле найден пустой размер")
    if isinstance(value, bool):
        raise ValueError(f"Некорректный размер: {value!r}")
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        if number == number.to_integral_value():
            return str(int(number))
        return format(number.normalize(), "f")
    result = str(value).strip().replace(",", ".")
    try:
        number = Decimal(result)
        if number == number.to_integral_value():
            return str(int(number))
        return format(number.normalize(), "f")
    except InvalidOperation:
        if not result:
            raise ValueError("В файле найден пустой размер")
        return result


def decimal_value(value: Any, field: str) -> Decimal:
    if value is None or str(value).strip() == "":
        raise ValueError(f"Пустое значение в поле «{field}»")
    try:
        result = Decimal(str(value).replace(" ", "").replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"Некорректное число в поле «{field}»: {value!r}") from exc
    return result.quantize(Decimal("0.01"))


def integer_value(value: Any, field: str, *, minimum: int = 0) -> int:
    number = decimal_value(value, field)
    if number != number.to_integral_value():
        raise ValueError(f"Ожидалось целое число в поле «{field}»: {value!r}")
    result = int(number)
    if result < minimum:
        raise ValueError(f"Поле «{field}» должно быть не меньше {minimum}")
    return result


def read_xlsx(path: Path, expected_headers: list[str]) -> list[dict[str, Any]]:
    if not path.is_file():
        raise FileNotFoundError(f"Не найден файл импорта: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = [clean_text(value) for value in next(rows)]
        except StopIteration:
            raise ValueError(f"Файл пустой: {path.name}")
        missing = [header for header in expected_headers if header not in headers]
        if missing:
            raise ValueError(
                f"В файле {path.name} не найдены столбцы: {', '.join(missing)}. "
                f"Обнаружены: {headers}"
            )
        indexes = {header: headers.index(header) for header in expected_headers}
        result: list[dict[str, Any]] = []
        for line_number, row in enumerate(rows, start=2):
            if not row or all(value is None for value in row):
                continue
            record = {
                header: row[index] if index < len(row) else None
                for header, index in indexes.items()
            }
            record["__line__"] = line_number
            result.append(record)
        return result
    finally:
        workbook.close()


def full_name_key(last_name: Any, first_name: Any, middle_name: Any) -> str:
    parts = [clean_text(part) for part in (last_name, first_name, middle_name)]
    return " ".join(part for part in parts if part).casefold()


def clear_existing_data(repositories: dict[str, Any]) -> None:
    """Remove imported entities through repositories, child records first."""
    delete_order = (
        ("order_items", "order_item"),
        ("orders", "order"),
        ("stock_items", "stock_item"),
        ("products", "product"),
        ("subcategories", "subcategory"),
        ("categories", "category"),
        ("sizes", "size"),
        ("users", "user"),
        ("roles", "role"),
    )
    for table_name, repository_key in delete_order:
        repository = repositories[repository_key]
        for entity in repository.get_all():
            repository.delete(entity)


def import_data(input_dir: Path, reset: bool = False) -> dict[str, int]:
    users_rows = read_xlsx(
        input_dir / FILES["users"], ["Фамилия", "Имя", "Отчетсво", "Логин", "Роль"]
    )
    products_rows = read_xlsx(
        input_dir / FILES["products"],
        [
            "Категория",
            "Подкатегория",
            "Изображение",
            "Наименование товара",
            "Производство",
            "Описание",
            "Состав",
            "Цена",
        ],
    )
    sizes_rows = read_xlsx(input_dir / FILES["sizes"], ["Размер"])
    stock_rows = read_xlsx(
        input_dir / FILES["stock"],
        [
            "Наименование товара",
            "Производство",
            "Размер",
            "Количество доступное для заказа",
        ],
    )
    orders_rows = read_xlsx(
        input_dir / FILES["orders"],
        [
            "Номер заказа",
            "Дата заказа",
            "ФИО",
            "Категория",
            "Наименование товара",
            "Производство",
            "Размер",
            "Количество",
            "Цена за единицу",
        ],
    )

    for row in users_rows:
        for column in ("Фамилия", "Имя", "Логин", "Роль"):
            clean_text(
                row[column],
                required=True,
                field=f"Users_import.xlsx:{row['__line__']}:{column}",
            )
    user_logins = [normalize_key(row["Логин"]) for row in users_rows]
    if len(user_logins) != len(set(user_logins)):
        raise ValueError("В Users_import.xlsx есть повторяющиеся логины")

    product_source: dict[tuple[str, str], dict[str, Any]] = {}
    for row in products_rows:
        key = (
            normalize_key(row["Наименование товара"]),
            normalize_key(row["Производство"]),
        )
        if not all(key):
            raise ValueError(
                f"В Products_import.xlsx пустое название или производство, строка {row['__line__']}"
            )
        if key in product_source:
            raise ValueError(
                f"Повторяется товар в Products_import.xlsx: {row['Наименование товара']} / {row['Производство']}"
            )
        product_source[key] = row
        decimal_value(row["Цена"], "Цена")

    size_values = {normalize_size(row["Размер"]) for row in sizes_rows}
    if len(size_values) != len(sizes_rows):
        raise ValueError(
            "В Sizes_import.xlsx есть повторяющиеся размеры после нормализации"
        )

    order_product_details: dict[tuple[str, str], dict[str, Any]] = {}
    for row in orders_rows:
        key = (
            normalize_key(row["Наименование товара"]),
            normalize_key(row["Производство"]),
        )
        detail = order_product_details.setdefault(
            key,
            {
                "name": clean_text(row["Наименование товара"], required=True),
                "manufacturer": clean_text(row["Производство"], required=True),
                "category": clean_text(row["Категория"], required=True),
                "price": decimal_value(row["Цена за единицу"], "Цена за единицу"),
            },
        )
        if detail["category"] != clean_text(row["Категория"], required=True):
            raise ValueError(
                f"У товара {detail['name']} в истории заказов указаны разные категории"
            )

    product_aliases: dict[tuple[str, str], tuple[str, str]] = {}
    for order_key in order_product_details:
        if order_key in product_source:
            product_aliases[order_key] = order_key
            continue
        order_name, order_manufacturer = order_key
        candidates = [
            product_key
            for product_key in product_source
            if product_key[1] == order_manufacturer
            and (
                product_key[0].startswith(order_name)
                or order_name.startswith(product_key[0])
            )
        ]
        if len(candidates) == 1:
            product_aliases[order_key] = candidates[0]
        elif len(candidates) > 1:
            raise ValueError(
                "Неоднозначное совпадение товара из истории заказов с каталогом: "
                f"{order_product_details[order_key]['name']} / {order_product_details[order_key]['manufacturer']}. "
                "Уточни название в исходном Excel."
            )
        else:
            product_aliases[order_key] = order_key

    synthetic_keys = {
        canonical_key
        for canonical_key in product_aliases.values()
        if canonical_key not in product_source
    }
    for row in stock_rows:
        key = (
            normalize_key(row["Наименование товара"]),
            normalize_key(row["Производство"]),
        )
        if key not in product_source and key not in order_product_details:
            raise ValueError(
                "Товар из остатков отсутствует в каталоге и истории заказов: "
                f"{row['Наименование товара']} / {row['Производство']}"
            )
        if key not in product_source:
            canonical_key = product_aliases.get(key, key)
            if canonical_key not in product_source:
                synthetic_keys.add(canonical_key)
        normalize_size(row["Размер"])
        integer_value(
            row["Количество доступное для заказа"], "Количество доступное для заказа"
        )
    for row in orders_rows:
        if normalize_size(row["Размер"]) not in size_values:
            raise ValueError(
                f"Размер {row['Размер']} из истории заказов отсутствует в Sizes_import.xlsx"
            )
        integer_value(row["Количество"], "Количество", minimum=1)
        decimal_value(row["Цена за единицу"], "Цена за единицу")
        if not isinstance(row["Дата заказа"], (datetime, date)):
            raise ValueError(
                f"Некорректная дата заказа в строке {row['__line__']}: {row['Дата заказа']!r}"
            )

    with Session(engine) as session:
        repos: dict[str, Any] = {
            "role": RoleRepository(session),
            "user": UserRepository(session),
            "category": CategoryRepository(session),
            "subcategory": SubcategoryRepository(session),
            "product": ProductRepository(session),
            "size": SizeRepository(session),
            "stock_item": StockItemRepository(session),
            "order": OrderRepository(session),
            "order_item": OrderItemRepository(session),
        }
        counts = defaultdict(int)
        existing_counts = {
            "roles": len(repos["role"].get_all()),
            "users": len(repos["user"].get_all()),
            "categories": len(repos["category"].get_all()),
            "subcategories": len(repos["subcategory"].get_all()),
            "products": len(repos["product"].get_all()),
            "sizes": len(repos["size"].get_all()),
            "stock_items": len(repos["stock_item"].get_all()),
            "orders": len(repos["order"].get_all()),
            "order_items": len(repos["order_item"].get_all()),
        }
        if reset:
            clear_existing_data(repos)
        elif any(existing_counts.values()):
            raise RuntimeError(
                "В БД уже есть данные (количество записей по таблицам: "
                f"{existing_counts}). Чтобы заменить их данными из Excel, запусти "
                "команду с --reset. Без этого флага скрипт не изменяет существующие данные."
            )

        try:
            roles_by_name: dict[str, Role] = {}
            for row in users_rows:
                original_role = clean_text(row["Роль"], required=True)
                role_name = ROLE_NAMES.get(original_role.casefold(), original_role)
                role_key = role_name.casefold()
                role = roles_by_name.get(role_key) or repos["role"].get_by_name(
                    role_name
                )
                if role is None:
                    role = repos["role"].create(Role(role_name=role_name))
                    counts["roles"] += 1
                roles_by_name[role_key] = role
                user = User(
                    last_name=clean_text(row["Фамилия"], required=True),
                    first_name=clean_text(row["Имя"], required=True),
                    middle_name=clean_text(row["Отчетсво"]),
                    login=clean_text(row["Логин"], required=True),
                    role_id=role.role_id,
                )
                repos["user"].create(user)
                counts["users"] += 1

            categories_by_name: dict[str, Category] = {}
            subcategories_by_key: dict[tuple[str, str], Subcategory] = {}

            def ensure_category(name: str) -> Category:
                name = clean_text(name, required=True)
                key = normalized_name(name)
                if key not in categories_by_name:
                    category = repos["category"].get_by_name(name)
                    if category is None:
                        category = repos["category"].create(Category(name=name))
                        counts["categories"] += 1
                    categories_by_name[key] = category
                return categories_by_name[key]

            def ensure_subcategory(
                category: Category, name: str | None
            ) -> Subcategory | None:
                if not name:
                    return None
                name = clean_text(name)
                if not name:
                    return None
                key = (normalized_name(category.name), normalized_name(name))
                if key not in subcategories_by_key:
                    subcategory = next(
                        (
                            item
                            for item in repos["subcategory"].get_by_category_id(
                                category.category_id
                            )
                            if normalized_name(item.name) == normalized_name(name)
                        ),
                        None,
                    )
                    if subcategory is None:
                        subcategory = repos["subcategory"].create(
                            Subcategory(name=name, category_id=category.category_id)
                        )
                        counts["subcategories"] += 1
                    subcategories_by_key[key] = subcategory
                return subcategories_by_key[key]

            product_map: dict[tuple[str, str], Product] = {}
            for row in products_rows:
                category = ensure_category(clean_text(row["Категория"], required=True))
                subcategory = ensure_subcategory(
                    category, clean_text(row["Подкатегория"])
                )
                product = Product(
                    category_id=category.category_id,
                    subcategory_id=subcategory.subcategory_id if subcategory else None,
                    image_path=clean_text(row["Изображение"]),
                    product_name=clean_text(row["Наименование товара"], required=True),
                    manufacturer=clean_text(row["Производство"], required=True),
                    description=clean_text(row["Описание"]),
                    composition=clean_text(row["Состав"]),
                    price=decimal_value(row["Цена"], "Цена"),
                )
                repos["product"].create(product)
                product_map[
                    (
                        normalize_key(product.product_name),
                        normalize_key(product.manufacturer),
                    )
                ] = product
                counts["products"] += 1

            for key in sorted(synthetic_keys):
                detail = order_product_details[key]
                category = ensure_category(detail["category"])
                product = Product(
                    category_id=category.category_id,
                    subcategory_id=None,
                    image_path=None,
                    product_name=detail["name"],
                    manufacturer=detail["manufacturer"],
                    description="Карточка восстановлена из истории заказов; описание отсутствует в файле импорта.",
                    composition=None,
                    price=detail["price"],
                )
                repos["product"].create(product)
                product_map[key] = product
                counts["products"] += 1
                counts["products_recovered_from_orders"] += 1

            for source_key, canonical_key in product_aliases.items():
                if source_key != canonical_key:
                    product_map[source_key] = product_map[canonical_key]
                    counts["product_name_aliases"] += 1

            sizes_by_value: dict[str, Size] = {}
            for row in sizes_rows:
                value = normalize_size(row["Размер"])
                size = repos["size"].create(Size(size_value=value))
                sizes_by_value[value] = size
                counts["sizes"] += 1

            for row in stock_rows:
                key = (
                    normalize_key(row["Наименование товара"]),
                    normalize_key(row["Производство"]),
                )
                product = product_map[key]
                size = sizes_by_value[normalize_size(row["Размер"])]
                stock = StockItem(
                    product_id=product.product_id,
                    size_id=size.size_id,
                    quantity_available=integer_value(
                        row["Количество доступное для заказа"],
                        "Количество доступное для заказа",
                    ),
                )
                repos["stock_item"].create(stock)
                counts["stock_items"] += 1

            users_by_full_name: dict[str, User] = {}
            for user in repos["user"].get_all():
                users_by_full_name[
                    full_name_key(user.last_name, user.first_name, user.middle_name)
                ] = user

            grouped_orders: dict[int, list[dict[str, Any]]] = defaultdict(list)
            for row in orders_rows:
                grouped_orders[
                    integer_value(row["Номер заказа"], "Номер заказа", minimum=1)
                ].append(row)

            for source_order_number in sorted(grouped_orders):
                lines = grouped_orders[source_order_number]
                first = lines[0]
                customer_name = clean_text(first["ФИО"], required=True)
                user = users_by_full_name.get(customer_name.casefold())
                if user is None:
                    raise ValueError(
                        f"Клиент «{customer_name}» из заказа {source_order_number} отсутствует в Users_import.xlsx"
                    )
                dates = {
                    (
                        row["Дата заказа"].date()
                        if isinstance(row["Дата заказа"], datetime)
                        else row["Дата заказа"]
                    )
                    for row in lines
                }
                names = {
                    clean_text(row["ФИО"], required=True).casefold() for row in lines
                }
                if len(dates) != 1 or len(names) != 1:
                    raise ValueError(
                        f"У строк заказа №{source_order_number} различаются дата или клиент"
                    )
                raw_date = first["Дата заказа"]
                order_date = (
                    raw_date
                    if isinstance(raw_date, datetime)
                    else datetime.combine(raw_date, datetime.min.time())
                )
                order = repos["order"].create(
                    Order(user_id=user.user_id, order_date=order_date)
                )
                counts["orders"] += 1

                for row in lines:
                    key = (
                        normalize_key(row["Наименование товара"]),
                        normalize_key(row["Производство"]),
                    )
                    product = product_map[product_aliases[key]]
                    size = sizes_by_value[normalize_size(row["Размер"])]
                    item = OrderItem(
                        order_id=order.order_id,
                        product_id=product.product_id,
                        size_id=size.size_id,
                        quantity=integer_value(
                            row["Количество"], "Количество", minimum=1
                        ),
                        price=decimal_value(row["Цена за единицу"], "Цена за единицу"),
                    )
                    repos["order_item"].create(item)
                    counts["order_items"] += 1

            session.commit()
        except Exception:
            session.rollback()
            raise

    return dict(counts)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Заполнить БД «Чудо Обувь» из XLSX-файлов"
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=DEFAULT_INPUT_DIR,
        help="Папка с пятью XLSX-файлами импорта",
    )
    parser.add_argument(
        "--reset",
        action="store_true",
        help="Удалить текущие данные приложения перед импортом",
    )
    args = parser.parse_args()

    try:
        create_tables()
        counts = import_data(args.input_dir.resolve(), reset=args.reset)
    except Exception as exc:
        print(f"Ошибка импорта: {exc}", file=sys.stderr)
        return 1

    print("Импорт завершён успешно:")
    for name, count in counts.items():
        print(f"  {name}: {count}")
    print(
        "Исторические заказы импортированы без повторного уменьшения складских остатков."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
